from __future__ import annotations

import io

import pytest
from openpyxl import Workbook, load_workbook

from translation_backend.app.application.project.translation_task.schemas import (
    TableColumn,
    TranslationTaskCreateRequest,
)
from translation_backend.app.infrastructure.spreadsheet.csv_adapter import CsvSpreadsheetAdapter
from translation_backend.app.infrastructure.spreadsheet.contracts import SpreadsheetSpec
from translation_backend.app.infrastructure.spreadsheet.xlsx_adapter import XlsxSpreadsheetAdapter


def test_translation_task_request_requires_model_and_unique_target_columns():
    with pytest.raises(ValueError):
        TranslationTaskCreateRequest(
            name="batch",
            source_language="zh-CN",
            target_languages=["en", "ko"],
            file_ids=[],
            api_key_id="00000000-0000-0000-0000-000000000001",
            model_type="openai.chat",
            model="gpt-test",
            source_column="source",
            target_columns=[TableColumn(language="en", column="en"), TableColumn(language="en", column="en-2")],
        )


def test_csv_adapter_reads_and_writes_only_explicit_translation_cells():
    content = "source,en,ko\n你好,Hello,\n再见,,Bye\n".encode()
    spec = SpreadsheetSpec(
        source_column="source",
        target_columns=[TableColumn(language="en", column="en"), TableColumn(language="ko", column="ko")],
        overwrite=True,
    )
    adapter = CsvSpreadsheetAdapter()
    cells = adapter.read(content, spec)
    assert [cell.source_text for cell in cells] == ["你好", "你好", "再见", "再见"]
    output = adapter.write(content, spec, {(2, 0, 1): "Hello new", (3, 0, 2): "再见韩文"})
    assert output.decode() == "source,en,ko\n你好,Hello new,\n再见,,再见韩文\n"


def test_xlsx_adapter_preserves_other_sheet_and_cell_style():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Translations"
    sheet.append(["source", "en"])
    sheet.append(["你好", ""])
    sheet["A2"].font = sheet["A1"].font.copy(bold=True)
    other = workbook.create_sheet("Other")
    other["A1"] = "keep"
    stream = io.BytesIO()
    workbook.save(stream)
    spec = SpreadsheetSpec(source_column="source", target_columns=[TableColumn(language="en", column="en")])
    adapter = XlsxSpreadsheetAdapter()
    cells = adapter.read(stream.getvalue(), spec)
    output = adapter.write(stream.getvalue(), spec, {("Translations", "A2", "B2"): "Hello"})
    result = load_workbook(io.BytesIO(output))
    assert cells[0].source_text == "你好"
    assert result["Translations"]["B2"].value == "Hello"
    assert result["Other"]["A1"].value == "keep"
    assert result["Translations"]["A2"].font.bold is True
