from __future__ import annotations

import io

from openpyxl import load_workbook

from .contracts import CellAddress, SpreadsheetCell, SpreadsheetSpec


class XlsxSpreadsheetError(ValueError):
    pass


class XlsxSpreadsheetAdapter:
    def read(self, content: bytes, spec: SpreadsheetSpec) -> list[SpreadsheetCell]:
        workbook = load_workbook(io.BytesIO(content), data_only=False)
        cells: list[SpreadsheetCell] = []
        for sheet in self._sheets(workbook, spec):
            try:
                headers = self._headers(sheet, spec)
            except XlsxSpreadsheetError:
                if spec.sheet_names:
                    raise
                continue
            for row in range(2, sheet.max_row + 1):
                source = sheet.cell(row, headers[spec.source_column]).value
                if not isinstance(source, str) or not source.strip() or source.startswith("="):
                    continue
                for target in spec.target_columns:
                    target_cell = sheet.cell(row, headers[target.column])
                    cells.append(SpreadsheetCell((sheet.title, f"{sheet.cell(row, headers[spec.source_column]).column_letter}{row}", target_cell.coordinate), source.strip(), target.language, target_cell.value if isinstance(target_cell.value, str) and target_cell.value.strip() else None))
        return cells

    def write(self, content: bytes, spec: SpreadsheetSpec, translations: dict[CellAddress, str]) -> bytes:
        workbook = load_workbook(io.BytesIO(content), data_only=False)
        for sheet in self._sheets(workbook, spec):
            try:
                self._headers(sheet, spec, create_missing=True)
            except XlsxSpreadsheetError:
                if spec.sheet_names:
                    raise
        for address, translated in translations.items():
            if len(address) != 3:
                continue
            sheet_name, _, target_coordinate = address
            if sheet_name not in workbook.sheetnames:
                continue
            cell = workbook[sheet_name][target_coordinate]
            if spec.overwrite or cell.value in (None, ""):
                cell.value = translated
        stream = io.BytesIO()
        workbook.save(stream)
        return stream.getvalue()

    def _sheets(self, workbook, spec: SpreadsheetSpec):
        if not spec.sheet_names:
            return workbook.worksheets
        missing = [name for name in spec.sheet_names if name not in workbook.sheetnames]
        if missing:
            raise XlsxSpreadsheetError(f"工作表不存在: {missing[0]}")
        return [workbook[name] for name in spec.sheet_names]

    @staticmethod
    def _headers(sheet, spec: SpreadsheetSpec, *, create_missing: bool = False) -> dict[str, int]:
        values = [sheet.cell(1, column).value for column in range(1, sheet.max_column + 1)]
        result: dict[str, int] = {}
        for name in [spec.source_column, *(item.column for item in spec.target_columns)]:
            matches = [index + 1 for index, value in enumerate(values) if value == name]
            if len(matches) > 1:
                raise XlsxSpreadsheetError(f"工作表列重复: {name}")
            if not matches:
                if not create_missing or name == spec.source_column:
                    raise XlsxSpreadsheetError(f"工作表缺少列: {name}")
                result[name] = sheet.max_column + 1
                sheet.cell(1, result[name]).value = name
            else:
                result[name] = matches[0]
        return result


__all__ = ["XlsxSpreadsheetAdapter", "XlsxSpreadsheetError"]
